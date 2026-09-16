import docx
import json

def extract_text_from_docx(file_path):
    try:
        doc = docx.Document(file_path)
        text = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Extract from tables as well
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                text.append(" | ".join(row_data))
                
        return "\n".join(text)
    except Exception as e:
        return str(e)

file1 = r"C:\Users\tmalu\Documents\Daily DSE Wrap 11 September 2026.docx"
file2 = r"C:\Users\tmalu\Documents\Current Prices 11 September 2026.docx"

text1 = extract_text_from_docx(file1)
text2 = extract_text_from_docx(file2)

with open('daily_wrap.txt', 'w', encoding='utf-8') as f:
    f.write(text1)

with open('current_prices.txt', 'w', encoding='utf-8') as f:
    f.write(text2)
