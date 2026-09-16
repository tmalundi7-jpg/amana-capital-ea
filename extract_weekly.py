import docx
import os

files = [
    r"C:\Users\tmalu\Documents\Daily DSE Wrap 07 September 2026.docx",
    r"C:\Users\tmalu\Documents\Daily DSE Wrap 08 September 2026.docx",
    r"C:\Users\tmalu\Documents\Daily DSE Wrap 09 September 2026.docx",
    r"C:\Users\tmalu\Documents\Daily DSE Wrap 10 September 2026.docx",
    r"C:\Users\tmalu\Documents\Daily DSE Wrap 11 September 2026.docx",
    r"C:\Users\tmalu\Documents\Current Prices 07 September 2026.docx",
    r"C:\Users\tmalu\Documents\Current Prices 08 September 2026.docx",
    r"C:\Users\tmalu\Documents\Current Prices 09 September 2026.docx",
    r"C:\Users\tmalu\Documents\Current Prices 10 September 2026.docx",
    r"C:\Users\tmalu\Documents\Current Prices 11 September 2026.docx"
]

with open("weekly_data.txt", "w", encoding="utf-8") as out:
    for file_path in files:
        out.write(f"\n\n--- FILE: {os.path.basename(file_path)} ---\n")
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    out.write(text + "\n")
            
            for table in doc.tables:
                out.write("\nTABLE:\n")
                for row in table.rows:
                    row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    out.write(" | ".join(row_data) + "\n")
        except Exception as e:
            out.write(f"ERROR reading {file_path}: {e}\n")
