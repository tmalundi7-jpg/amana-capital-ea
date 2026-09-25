import sys
from docx import Document

with open('extracted_wrap.txt', 'w', encoding='utf-8') as f:
    doc = Document(r'C:\Users\tmalu\Documents\Daily DSE Wrap 25 September 2026.docx')
    for p in doc.paragraphs:
        f.write(p.text + '\n')
    f.write("\nTABLES:\n")
    for table in doc.tables:
        for row in table.rows:
            f.write(" | ".join([c.text.strip().replace('\n', ' ') for c in row.cells]) + '\n')
