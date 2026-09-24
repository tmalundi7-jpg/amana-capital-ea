from docx import Document

PIPE = " | "

for label, path in [
    ("DAILY WRAP", r"C:\Users\tmalu\Documents\Daily DSE Wrap 23 September 2026.docx"),
    ("CURRENT PRICES", r"C:\Users\tmalu\Documents\Current Prices 23 September 2026.docx"),
]:
    print(f"\n\n========== {label} ==========")
    try:
        doc = Document(path)
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                print(f"[{i}] {para.text}")
        for t_idx, table in enumerate(doc.tables):
            print(f"\n=== TABLE {t_idx} ===")
            for r_idx, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells]
                print(f"  Row {r_idx}: {PIPE.join(cells)}")
    except Exception as e:
        print(f"ERROR: {e}")
