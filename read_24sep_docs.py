import docx
import sys
import os

def extract_docx(filepath):
    if not os.path.exists(filepath):
        print(f"ERROR: File not found: {filepath}")
        return
    
    print(f"\n--- EXTRACTING: {os.path.basename(filepath)} ---")
    doc = docx.Document(filepath)
    
    print("\nPARAGRAPHS:")
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            print(f"[{i}] {text}")
            
    print("\nTABLES:")
    for t_idx, table in enumerate(doc.tables):
        print(f"\n=== TABLE {t_idx} ===")
        for r_idx, row in enumerate(table.rows):
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"Row {r_idx}: " + " | ".join(cells))

# Paths
wrap_path = r"C:\Users\tmalu\Documents\Daily DSE Wrap 24 September 2026.docx"
prices_path = r"C:\Users\tmalu\Documents\Current Prices 24 September 2026.docx"

extract_docx(wrap_path)
extract_docx(prices_path)
