import docx
import os

docs = [
    r"C:\Users\tmalu\Documents\Daily DSE Wrap 25 September 2026.docx",
    r"C:\Users\tmalu\Documents\Current Prices 25 September 2026.docx"
]

out_file = r"C:\Users\tmalu\.gemini\antigravity\scratch\Amana-capital-ea\extracted_25sep_data.txt"

with open(out_file, "w", encoding="utf-8") as f:
    for d in docs:
        if not os.path.exists(d):
            f.write(f"MISSING: {d}\n\n")
            continue
            
        f.write(f"=== CONTENT OF {os.path.basename(d)} ===\n")
        doc = docx.Document(d)
        f.write("PARAGRAPHS:\n")
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                f.write(f"[{i}] {p.text.strip()}\n")
                
        f.write("\nTABLES:\n")
        for t_idx, table in enumerate(doc.tables):
            f.write(f"--- Table {t_idx} ---\n")
            for r_idx, row in enumerate(table.rows):
                row_data = [cell.text.strip() for cell in row.cells]
                f.write(f"Row {r_idx}: " + " | ".join(row_data) + "\n")
        f.write("\n" + "="*50 + "\n\n")

print(f"Extraction complete. Check {out_file}")
