import docx
import re
import shutil
import json

wrap_path = r"C:\Users\tmalu\Documents\Daily DSE Wrap 11 September 2026.docx"
prices_path = r"C:\Users\tmalu\Documents\Current Prices 11 September 2026.docx"

def extract_wrap(path):
    doc = docx.Document(path)
    text = []
    for p in doc.paragraphs:
        if p.text.strip():
            text.append(p.text.strip())
    for t in doc.tables:
        for r in t.rows:
            row_data = [c.text.strip().replace('\n', ' ') for c in r.cells]
            text.append(' | '.join(row_data))
    return '\n'.join(text)

def extract_prices(path):
    doc = docx.Document(path)
    text = []
    for t in doc.tables:
        for r in t.rows:
            row_data = [c.text.strip().replace('\n', ' ') for c in r.cells]
            text.append(' | '.join(row_data))
    return '\n'.join(text)

wrap_text = extract_wrap(wrap_path)
prices_text = extract_prices(prices_path)

with open('output_wrap.txt', 'w', encoding='utf-8') as f:
    f.write(wrap_text)
    
with open('output_prices.txt', 'w', encoding='utf-8') as f:
    f.write(prices_text)

print("Extraction completed!")
